"""
Import Trendyol product export (urunlerimiz.xlsx) into Firestore.

Usage:
  cd backend
  python scripts/import_trendyol_excel.py path/to/urunlerimiz.xlsx
  python scripts/import_trendyol_excel.py path/to/file.xlsx --dry-run
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from app.core.firebase import get_db
from app.core.utils import slugify

IMAGE_PREFIX = "Görsel"


def _find_header(headers: tuple, *needles: str) -> str | None:
    for h in headers:
        if h is None:
            continue
        s = str(h)
        if all(n in s for n in needles):
            return str(h)
    return None


def _find_exact(headers: tuple, name: str) -> str | None:
    for h in headers:
        if h is not None and str(h).strip() == name:
            return str(h)
    return None


def _parse_price_try(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(round(float(value) * 100))
    s = str(value).strip().replace(",", ".")
    s = re.sub(r"[^\d.]", "", s)
    if not s:
        return 0
    return int(round(float(s) * 100))


def _parse_int(value: Any, default: int = 0) -> int:
    if value is None:
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _collect_images(row: dict[str, Any], headers: tuple) -> list[str]:
    urls: list[str] = []
    for h in headers:
        if h and str(h).startswith(IMAGE_PREFIX):
            url = row.get(str(h))
            if url and str(url).startswith("http"):
                urls.append(str(url).strip())
    return urls


def _infer_type(name: str, category: str) -> str:
    name_l = name.lower()
    cat_l = category.lower()
    if any(k in name_l for k in ("hediye", "koleksiyon", "paketi", "kutu")):
        return "gift_box"
    if "set" in cat_l or "seti" in cat_l or " set" in name_l:
        return "curation_set"
    if any(k in name_l for k in ("seti", "set ", "'lü", "parça", "parca", "avantaj paketi")):
        return "curation_set"
    return "single"


def _row_active(row: dict[str, Any], durum_key: str | None) -> bool:
    if not durum_key:
        return True
    durum = row.get(durum_key)
    if durum is None or str(durum).strip() in ("", "None"):
        return True
    return False


def load_rows(path: str) -> tuple[tuple, list[dict[str, Any]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = None
    for name in wb.sheetnames:
        sheet = wb[name]
        if sheet.max_row and sheet.max_row > 10:
            ws = sheet
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("Empty spreadsheet")

    headers = tuple(rows[0])
    data = []
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        data.append(dict(zip(headers, row)))
    return headers, data


def build_products(headers: tuple, rows: list[dict[str, Any]]) -> tuple[dict[str, dict], dict[str, str]]:
    model_key = _find_header(headers, "Model", "Kodu") or _find_header(headers, "Model")
    name_key = _find_header(headers, "rün", "Ad") or _find_header(headers, "Ad")
    cat_key = _find_header(headers, "Kategori")
    price_key = _find_header(headers, "Sat", "Fiyat") or _find_header(headers, "Satilacak")
    stock_key = _find_header(headers, "Stok")
    barkod_key = _find_exact(headers, "Barkod") or _find_header(headers, "Barkod")
    color_key = _find_header(headers, "Rengi")
    size_key = _find_exact(headers, "Beden") or _find_header(headers, "Beden")
    desc_key = _find_header(headers, "Açıklama") or _find_header(headers, "Aciklama")
    durum_key = _find_exact(headers, "Durum")

    if not model_key or not name_key or not price_key:
        raise ValueError(f"Missing required columns. Found headers: {headers[:15]}")

    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        model = str(row.get(model_key) or row.get(barkod_key) or "").strip()
        if not model:
            continue
        groups[model].append(row)

    categories: dict[str, str] = {}
    products: dict[str, dict] = {}
    now = datetime.now(timezone.utc)

    for model_kodu, group_rows in groups.items():
        first = group_rows[0]
        name_tr = str(first.get(name_key) or model_kodu).strip()
        category_name = str(first.get(cat_key) or "Genel").strip()
        cat_slug = slugify(category_name) or "genel"
        categories.setdefault(category_name, cat_slug)

        description = first.get(desc_key)
        description_tr = str(description).strip() if description else ""

        variants = []
        all_images: list[str] = []
        any_active = False
        min_price = None

        for row in group_rows:
            sku = str(row.get(barkod_key) or model_kodu).strip()
            price_try = _parse_price_try(row.get(price_key))
            stock = _parse_int(row.get(stock_key), 0)
            active_row = _row_active(row, durum_key)

            options: dict[str, str] = {}
            if color_key and row.get(color_key):
                options["renk"] = str(row.get(color_key)).strip()
            if size_key and row.get(size_key):
                options["beden"] = str(row.get(size_key)).strip()

            images = _collect_images(row, headers)
            all_images.extend(images)

            if active_row:
                any_active = True

            if min_price is None or (price_try > 0 and price_try < min_price):
                min_price = price_try

            variants.append(
                {
                    "sku": sku,
                    "options": options,
                    "price_try": price_try,
                    "stock": max(stock, 0),
                    "low_stock_threshold": 5,
                }
            )

        # dedupe images preserve order
        seen: set[str] = set()
        unique_images: list[str] = []
        for url in all_images:
            if url not in seen:
                seen.add(url)
                unique_images.append(url)

        product_type = _infer_type(name_tr, category_name)
        base_price = min_price or (variants[0]["price_try"] if variants else 0)
        slug = slugify(model_kodu) or slugify(str(first.get(barkod_key) or name_tr))
        if not slug:
            slug = re.sub(r"[^a-z0-9-]", "", model_kodu.lower())

        products[slug] = {
            "name_tr": name_tr,
            "slug": slug,
            "type": product_type,
            "description_tr": description_tr,
            "images": unique_images[:8],
            "base_price_try": base_price,
            "category_slug": cat_slug,
            "category_name": category_name,
            "variants": variants,
            "is_featured": product_type != "single" and any_active,
            "gift_wrap_available": product_type != "single",
            "active": any_active,
            "source_model_kodu": model_kodu,
            "updated_at": now,
        }

    return products, categories


def import_to_firestore(
    products: dict[str, dict],
    categories: dict[str, str],
    *,
    dry_run: bool = False,
    clear_existing: bool = False,
) -> None:
    if dry_run:
        print(f"[dry-run] Would import {len(categories)} categories, {len(products)} products")
        for slug, p in list(products.items())[:3]:
            print(f"  - {p['name_tr'][:60]} | {len(p['variants'])} variants | TRY {p['base_price_try']/100:.2f}")
        return

    db = get_db()
    now = datetime.now(timezone.utc)

    try:
        # Verify Firestore is provisioned
        db.collection("_health").document("ping").set({"ok": True}, merge=True)
    except Exception as exc:
        msg = str(exc)
        if "does not exist" in msg or "NOT_FOUND" in msg:
            raise SystemExit(
                "Firestore database is not created yet.\n"
                "Firebase Console → Firestore Database → Create database "
                "(region: europe-west1)\n"
                "https://console.firebase.google.com/project/dami-beauty-353b0/firestore"
            ) from exc
        raise

    if clear_existing:
        print("Clearing existing products...")
        for doc in db.collection("products").stream():
            doc.reference.delete()

    cat_ids: dict[str, str] = {}
    sort_order = 0
    for name, cat_slug in sorted(categories.items()):
        payload = {
            "name_tr": name,
            "slug": cat_slug,
            "sort_order": sort_order,
            "active": True,
        }
        ref = db.collection("categories").document(cat_slug)
        ref.set(payload, merge=True)
        cat_ids[cat_slug] = cat_slug
        sort_order += 1

    created, updated = 0, 0
    batch_count = 0
    batch = db.batch()

    for slug, product in products.items():
        cat_slug = product.pop("category_slug")
        product.pop("category_name", None)
        product.pop("source_model_kodu", None)
        product["category_id"] = cat_ids.get(cat_slug)
        product["campaign_id"] = None

        ref = db.collection("products").document(slug)
        snap = ref.get()
        if snap.exists:
            product["created_at"] = snap.to_dict().get("created_at") or now
            batch.set(ref, product, merge=True)
            updated += 1
        else:
            product["created_at"] = now
            batch.set(ref, product)
            created += 1

        batch_count += 1
        if batch_count >= 400:
            batch.commit()
            batch = db.batch()
            batch_count = 0

    if batch_count:
        batch.commit()

    print(f"Done: {created} created, {updated} updated, {len(cat_ids)} categories.")


def main():
    parser = argparse.ArgumentParser(description="Import Trendyol Excel into Firestore")
    parser.add_argument("xlsx_path", help="Path to urunlerimiz.xlsx")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument(
        "--clear-existing",
        action="store_true",
        help="Delete all products before import (keeps categories updated)",
    )
    args = parser.parse_args()

    if not os.path.isfile(args.xlsx_path):
        print(f"File not found: {args.xlsx_path}", file=sys.stderr)
        sys.exit(1)

    print(f"Reading {args.xlsx_path}...")
    headers, rows = load_rows(args.xlsx_path)
    print(f"Loaded {len(rows)} rows")
    products, categories = build_products(headers, rows)
    print(f"Grouped into {len(products)} products, {len(categories)} categories")
    import_to_firestore(
        products,
        categories,
        dry_run=args.dry_run,
        clear_existing=args.clear_existing,
    )


if __name__ == "__main__":
    main()
